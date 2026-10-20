# -*- coding: utf-8 -*-
"""Generate docs/IDS_Universal_Guide.docx"""
from docx import Document
from docx.shared import Pt, RGBColor
from pathlib import Path

OUT = Path(r"E:\CSE499B\Latest One(FL+RL)\docs\IDS_Universal_Guide.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(11)


def h1(t): doc.add_heading(t, 1)
def h2(t): doc.add_heading(t, 2)
def p(t, bold=False):
    para = doc.add_paragraph()
    r = para.add_run(t); r.bold = bold
def bullet(t): doc.add_paragraph(t, style='List Bullet')
def num(t): doc.add_paragraph(t, style='List Number')
def code(t):
    para = doc.add_paragraph()
    r = para.add_run(t)
    r.font.name = 'Consolas'; r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

doc.add_heading('SDN-IoT IDS — Universal Guide (FL + RL Combined)', 0)
sub = doc.add_paragraph('Latest One (FL+RL) package — how Federated Learning and Reinforcement Learning were executed, combined, and how to improve the system further.')
sub.runs[0].italic = True

h1('1. What this package is')
p('Folder: E:\\CSE499B\\Latest One(FL+RL)')
bullet('IDS_FL_RL_Combined.ipynb — full pipeline (Sections 1–8 training identical to IDS_Advanced_GPU; Sections 9–10 = real FL + 5 RL agents)')
bullet('Saved model/ — TM, DL, RL, FL checkpoints + scaler / label encoder / feature columns')
bullet('federated/ — 1 Flower server + 3 clients (real multi-process FedAvg)')
bullet('ids_server.py + templates/dashboard.html — localhost UI with best-model-only prediction and live FL/RL proof')
bullet('docs/IDS_Universal_Guide.docx — this document')

h1('2. How the parts were combined')
h2('2.1 Base training (unchanged)')
p('Sections 1–8 of IDS_Advanced_GPU.ipynb were kept: cleaning, SMOTE, 12 traditional models, 8 DL models, HybridIDSNet, comparison dashboard. Artifacts now save under Saved model/ instead of Cursor/.')

h2('2.2 Federated Learning (real, not simulated)')
p('Section 9 partitions the balanced training set into 3 non-IID client shards (Dirichlet) plus a shared test.npz.')
p('Training uses Flower FedAvg with:')
bullet('1 main server process (server.py on port 8080)')
bullet('3 client processes (client.py 0/1/2), each with private data')
bullet('Architecture: ResNet-MLP — best deep model from the comparison that supports weight averaging')
p('Why not XGBoost for FL? XGBoost had the best overall F1, but tree ensembles do not participate in classic FedAvg (no shared differentiable weights). ResNet-MLP is the best federatable backbone. The held-out test set is evaluated on the server every round.')
p('Live evidence written during a run:')
bullet('federated/status.json — server round, accuracy, PID, timestamps')
bullet('federated/client_0_status.json … client_2_status.json — per-client heartbeats')
bullet('Saved model/FL/federated_fedavg.pt — global weights')

h2('2.3 Reinforcement Learning (from COMPLETE_RL_FULL_CODE)')
p('The old single Gymnasium DQN cell was removed. Section 10 now trains five agents:')
bullet('DQN-Basic — MLP Q-network')
bullet('DQN-Improved — Dueling architecture')
bullet('PPO — Actor-Critic policy gradient')
bullet('Rainbow-DQN — Dueling + Noisy Nets')
bullet('QR-DQN — distributional RL (51 quantiles)')
p('They use the same scaled feature space as supervised models, with cost-sensitive rewards (rare classes worth more). Checkpoints go to Saved model/RL/. This is real PyTorch training and inference — not a mock loop.')

h2('2.4 Web UI integration')
p('The UI no longer lets you pick among 23 models or shows a performance leaderboard.')
bullet('Attack prediction is locked to the best supervised model: XGBoost')
bullet('Every Predict click also runs the FL global model (if present) and all loaded RL agents')
bullet('The Live FL / RL Systems panel polls /live every 3 seconds so you can watch real Flower rounds and RL agent presence')
p('That dual-path inference is the UI proof that FL and RL are live systems, not decorative labels.')

h1('3. How to execute everything')
h2('3.1 One-time setup')
num('Open IDS_FL_RL_Combined.ipynb in the Latest One(FL+RL) folder')
num('Ensure ASEADOS_SDN_IoT.csv is reachable (parent E:\\CSE499B\\ or project root)')
num('Run Section 1 install cell (flwr, gymnasium optional, flask, torch CUDA build, etc.)')

h2('3.2 Train supervised models')
num('Run Sections 2–8 top to bottom (data → TM → DL → custom → comparison)')
num('Confirm Saved model/TM/xgboost.joblib exists after training')

h2('3.3 Run Federated Learning')
num('Run Section 9.1 (partition shards into federated/data/)')
num('Double-click federated/run_federated.bat  (or set RUN_FL=True in 9.3)')
num('Watch 4 terminals: FL-SERVER + FL-CLIENT-0/1/2')
num('When finished, run Section 9.4 to score and copy the global model into Saved model/FL/')
p('Detailed FL steps: federated/README.md')

h2('3.4 Run Reinforcement Learning')
num('Run Section 10 cells (trains all 5 agents and writes Saved model/RL/*.pt)')

h2('3.5 Save + launch UI')
num('Run Section 11 (scaler, label encoder, results_summary, model_info.json)')
num('In a terminal from Latest One(FL+RL):')
code('python ids_server.py')
num('Open http://127.0.0.1:7860')
num('Click Predict — compare Best model / RL agent / FL global boxes')
num('Watch Live FL / RL Systems while a federated run is active')

h1('4. How to verify it is NOT simulation-based')
h2('Federated Learning')
bullet('Four separate OS processes appear (Task Manager / four cmd windows)')
bullet('status.json round and updated_at change over time')
bullet('Client status files flip through connecting → local_training → fit_done')
bullet('Server prints global_test_acc each round')
bullet('UI heartbeat turns green and shows Server PID')

h2('Reinforcement Learning')
bullet('Section 10 prints per-agent training steps and test metrics')
bullet('Saved model/RL contains five .pt files')
bullet('UI lists loaded agents; Predict fills “RL agent says” with a real action + confidence')
bullet('/live returns mode: real_pytorch_inference')

h1('5. Design decisions (short)')
bullet('UI predictor = XGBoost (best F1 on the held-out test set in results_summary.csv)')
bullet('FL backbone = ResNet-MLP (best neural model compatible with FedAvg)')
bullet('RL = five modern agents from COMPLETE_RL_FULL_CODE, wired into RESULTS + Saved model/RL')
bullet('Performance leaderboard removed from the web UI (as requested)')

h1('6. What you can do next to make the system better')
h2('6.1 Federated Learning')
bullet('Try FedProx or SCAFFOLD if client drift is high on non-IID shards')
bullet('Add differential privacy (DP-SGD) on client updates for stronger privacy claims')
bullet('Increase FL_ROUNDS and local_epochs once GPU time allows')
bullet('Explore federated XGBoost / tree ensembles (Flower has experimental tree strategies) if you need the absolute best model inside FL')
bullet('Deploy clients on separate machines/VMs to demonstrate cross-host FL')

h2('6.2 Reinforcement Learning')
bullet('Raise EPOCHS and use a prioritized replay buffer for DQN-family agents')
bullet('Add a true multi-step Gymnasium environment (delayed rewards, episode-level false-alarm cost)')
bullet('Online fine-tuning: update RL agents from analyst feedback (human-in-the-loop)')
bullet('Ensemble voting between best RL agent and XGBoost with a calibrated threshold')

h2('6.3 System / product')
bullet('Stream live flows from an SDN controller or Zeek/CICFlowMeter instead of CSV paste')
bullet('Alerting webhooks (Slack/email) when Best + RL + FL agree on an attack')
bullet('Continual learning schedule: nightly FL rounds + weekly RL refresh')
bullet('Model cards + drift monitoring on feature distributions')
bullet('Hardening: auth on the dashboard, rate limits, audit logs')

h1('7. Quick file map')
code("""Latest One(FL+RL)/
  IDS_FL_RL_Combined.ipynb
  ids_server.py
  ids_models.py
  requirements.txt
  Saved model/
    TM/  DL/  RL/  FL/
    scaler.joblib  label_encoder.joblib  feature_cols.joblib
  federated/
    server.py  client.py  fl_model.py
    run_federated.bat  README.md  data/
  templates/dashboard.html
  docs/IDS_Universal_Guide.docx
""")

h1('8. Troubleshooting')
bullet('UI says best model missing → run Sections 5 + 11 so xgboost.joblib exists')
bullet('FL clients hang → start server first; FL_CLIENTS must be 3; firewall must allow 8080')
bullet('RL box shows “not loaded” → run Section 10 then restart ids_server.py')
bullet('CUDA false → install CUDA PyTorch wheel matching your GPU')
bullet('Dataset not found → place ASEADOS_SDN_IoT.csv in E:\\CSE499B\\ or the project folder')

doc.save(OUT)
print('Wrote', OUT)
